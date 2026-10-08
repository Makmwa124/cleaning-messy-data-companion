"""Rebuild every practice file used in Cleaning Messy Data.

Run from the companion folder:  python scripts/make_data.py
All companies, people, emails and phone numbers are fictional.
"""
import random
import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

rng = np.random.default_rng(2026)
random.seed(2026)
RAW = Path("data/raw")
REF = Path("data/reference")
RAW.mkdir(parents=True, exist_ok=True)
REF.mkdir(parents=True, exist_ok=True)

FIRST = ["James", "Olivia", "Jack", "Charlotte", "Noah", "Amelia",
         "William", "Isla", "Thomas", "Mia", "Lachlan", "Grace",
         "Henry", "Ava", "Oliver", "Chloe", "Liam", "Ruby", "Ethan",
         "Sophie", "Zoë", "José", "Siobhan", "Mateo", "Aisha",
         "Wei", "Priyanka", "Kiri", "Niamh", "Renée", "Hamish",
         "Tegan", "Brodie", "Jasmine", "Darcy", "Mohammed", "Linh",
         "Ngaio", "Callum", "Ellie"]
LAST = ["Smith", "Jones", "Williams", "Brown", "Wilson", "Taylor",
        "Nguyen", "Johnson", "Martin", "White", "Anderson", "Walker",
        "Thompson", "Harris", "Lee", "Ryan", "Kelly", "King",
        "O’Brien", "McDonald", "Müller", "Novak", "Papadopoulos",
        "Hughes", "Campbell", "Murray", "Stewart", "Clarke", "Young",
        "Mitchell", "Fraser", "Doyle", "Kowalski", "Patel", "Singh",
        "Rossi", "D’Angelo", "Gallagher", "Kaur", "Tran"]
TOWNS = [("Murray Bridge", "5253"), ("Mount Gambier", "5290"),
         ("Renmark", "5341"), ("Berri", "5343"), ("Loxton", "5333"),
         ("Waikerie", "5330"), ("Pinnaroo", "5304"),
         ("Lameroo", "5302"), ("Karoonda", "5307"),
         ("Naracoorte", "5271"), ("Bordertown", "5268"),
         ("Millicent", "5280"), ("Victor Harbor", "5211"),
         ("Mount Barker", "5251"), ("Tanunda", "5352"),
         ("Clare", "5453"), ("Port Lincoln", "5606"),
         ("Whyalla", "5600"), ("Port Augusta", "5700"),
         ("Gawler", "5118"), ("Darwin", "0800"),
         ("Alice Springs", "0870")]
TOWN_W = np.array([14, 9, 8, 7, 8, 5, 6, 5, 4, 4, 4, 3, 3, 4, 3,
                   2, 2, 2, 2, 3, 1.2, 0.8])
TOWN_W = TOWN_W / TOWN_W.sum()
TYPOS = {"Mount Gambier": ["Mt Gambier", "Mount Gambeir", "MOUNT GAMBIER"],
         "Murray Bridge": ["Murray Bdge", "Murrray Bridge", "murray bridge"],
         "Renmark": ["Renmark ", "Renmrk"],
         "Victor Harbor": ["Victor Harbour", "Victor Hbr"],
         "Mount Barker": ["Mt Barker", "Mount Baker"],
         "Port Lincoln": ["Pt Lincoln"],
         "Naracoorte": ["Naracorte"],
         "Loxton": ["Loxtn"]}
INTERESTS = ["Garden", "Kitchen", "Outdoor", "Tools", "Lighting", "Pets"]

# ---------------------------------------------------------------
# 1. Mallee Hardware customer export (the main messy file)
# ---------------------------------------------------------------
N = 1180
rows = []
for i in range(N):
    fn = random.choice(FIRST)
    ln = random.choice(LAST)
    town, pc = TOWNS[rng.choice(len(TOWNS), p=TOWN_W)]
    channel = random.choices(["In store", "Online", "Phone"],
                             [0.55, 0.38, 0.07])[0]
    join = pd.Timestamp("2018-07-01") + pd.Timedelta(
        days=int(rng.integers(0, 2900)))
    birth = pd.Timestamp("1945-01-01") + pd.Timedelta(
        days=int(rng.integers(0, 21000)))
    pts = int(rng.gamma(1.6, 900))
    picks = random.sample(INTERESTS, random.choice([0, 1, 1, 2, 2, 3]))
    mobile = f"04{rng.integers(10, 99)}{rng.integers(100000, 999999)}"
    rows.append(dict(cust_no=f"{i + 1:06d}", first=fn, last=ln,
                     town=town, postcode=pc, channel=channel,
                     join=join, birth=birth, points=pts, picks=picks,
                     mobile=mobile))

out = []
for r in rows:
    fn, ln = r["first"], r["last"]
    # name layout varies by who keyed it in
    style = random.random()
    if style < 0.55:
        name = f"{fn} {ln}"
    elif style < 0.75:
        name = f"{ln.upper()}, {fn}"
    elif style < 0.9:
        name = f"{fn.lower()} {ln.lower()}"
    else:
        name = f"{fn}  {ln} "          # double and trailing space
    if random.random() < 0.06:
        name = name.replace(" ", "\xa0", 1)   # non-breaking space
    email_user = f"{fn}.{ln}".lower().replace("’", "")
    email_user = (email_user.replace("ë", "e").replace("é", "e")
                  .replace("ü", "u").replace(" ", ""))
    dom = random.choice(["outlook.com", "gmail.com", "bigpond.com",
                         "internode.on.net", "yahoo.com.au"])
    email = f"{email_user}{random.randint(1, 99)}@{dom}"
    e = random.random()
    if e < 0.15:
        email = email.upper()
    elif e < 0.20:
        email = " " + email
    elif e < 0.23:
        email = email.replace(".com", "")       # missing domain ending
    elif e < 0.29:
        email = random.choice(["", "n/a", "none", "-"])
    m = r["mobile"]
    p = random.random()
    if p < 0.30:
        phone = f"{m[:4]} {m[4:7]} {m[7:]}"
    elif p < 0.50:
        phone = f"+61 {m[1:4]} {m[4:7]} {m[7:]}"
    elif p < 0.62:
        phone = m[1:]                            # lost leading zero
    elif p < 0.72:
        phone = f"61{m[1:]}"
    elif p < 0.82:
        phone = f"(08) 8{random.randint(100, 999)} {random.randint(1000, 9999)}"
    elif p < 0.92:
        phone = m
    else:
        phone = ""
    town = r["town"]
    if town in TYPOS and random.random() < 0.18:
        town = random.choice(TYPOS[town])
    state = "NT" if r["postcode"].startswith("0") else "SA"
    s = random.random()
    if state == "SA":
        if s < 0.62:
            st = "SA"
        elif s < 0.72:
            st = "S.A."
        elif s < 0.84:
            st = "South Australia"
        elif s < 0.92:
            st = "sa"
        elif s < 0.96:
            st = "SA "
        else:
            st = "Sth Aust"
    else:
        st = random.choice(["NT", "N.T.", "Northern Territory", "nt"])
    pc = r["postcode"].lstrip("0")               # Excel dropped zeros
    # join date: three formats from three systems
    j = r["join"]
    f = random.random()
    if j.year < 2021:
        jd = j.strftime("%d/%m/%Y")
        if random.random() < 0.3:
            jd = j.strftime("%-d/%-m/%Y")
    elif j.year < 2023:
        jd = j.strftime("%Y-%m-%d")
    else:
        jd = j.strftime("%d-%b-%y")
    b = r["birth"]
    if r["channel"] == "In store" and random.random() < 0.45:
        bd = ""
    elif random.random() < 0.05:
        bd = ""
    elif random.random() < 0.035:
        bd = "01/01/1900"
    else:
        bd = b.strftime("%d/%m/%Y")
    q = random.random()
    if q < 0.03:
        pts = "N/A"
    elif q < 0.05:
        pts = "-1"
    elif q < 0.06:
        pts = ""
    else:
        pts = f"{r['points']:,}"
    picks = r["picks"]
    sep = random.choice(["; ", ";", ", ", ","])
    interests = sep.join(x if random.random() < 0.8 else x.lower()
                         for x in picks)
    upd = j + pd.Timedelta(days=int(rng.integers(0, 400)))
    out.append({"cust_no": r["cust_no"], "full_name": name,
                "email": email, "phone": phone, "suburb": town,
                "postcode": pc, "state": st,
                "signup_channel": r["channel"], "join_date": jd,
                "birth_date": bd, "loyalty_points": pts,
                "interests": interests,
                "last_updated": upd.strftime("%Y-%m-%d")})

cust = pd.DataFrame(out)

# same person signed up twice (different cust_no, messier details)
twins = cust.sample(24, random_state=7).copy()
twins["cust_no"] = [f"{N + 1 + k:06d}" for k in range(len(twins))]
twins["email"] = twins["email"].str.upper().str.strip()
twins["phone"] = twins["phone"].str.replace(" ", "", regex=False)
twins["signup_channel"] = "Online"
twins["last_updated"] = "2026-05-02"
# same cust_no twice with conflicting details (a system re-sync)
clash = cust.sample(15, random_state=11).copy()
clash["suburb"] = [random.choice([t for t, _ in TOWNS[:15]])
                   for _ in range(len(clash))]
clash["last_updated"] = "2026-06-18"
# exact duplicates (the export overlapped by one page)
exact = cust.iloc[400:440].copy()
cust = pd.concat([cust, twins, clash, exact], ignore_index=True)
cust = cust.sample(frac=1, random_state=3).reset_index(drop=True)
with open(RAW / "mallee_customers.csv", "w", encoding="cp1252",
          newline="") as fh:
    cust.to_csv(fh, index=False)

# ---------------------------------------------------------------
# 2. Banksia customers (already clean) for matching in Chapter 9
# ---------------------------------------------------------------
bk = []
for i in range(3200):
    fn = random.choice(FIRST)
    ln = random.choice(LAST)
    town, pc = TOWNS[rng.choice(len(TOWNS), p=TOWN_W)]
    bk.append({"customer_id": f"B{10001 + i}", "first_name": fn,
               "last_name": ln,
               "email": f"{fn}.{ln}{random.randint(100, 999)}@example.com"
               .lower().replace("’", "").replace("ë", "e")
               .replace("é", "e").replace("ü", "u"),
               "mobile": f"04{rng.integers(10, 99)}"
                         f"{rng.integers(100000, 999999)}",
               "suburb": town, "postcode": pc,
               "state": "NT" if pc.startswith("0") else "SA"})
bkc = pd.DataFrame(bk)
# 38 Mallee customers already shop at Banksia (same email)
shared = cust.drop_duplicates("cust_no").sample(60, random_state=21)
shared = shared[shared["email"].str.contains("@", na=False)]
shared = shared[shared["email"].str.contains(r"\.com|\.net", na=False)]
shared = shared.head(38)
for k, (_, r) in enumerate(shared.iterrows()):
    bkc.loc[k * 50, "email"] = r["email"].strip().lower()

# ---------------------------------------------------------------
# 3. Web orders (UTC timestamps, units, sentinels, duplicates)
# ---------------------------------------------------------------
PRODUCTS = [("BK-1001", "Terracotta pot 30 cm", 24.95, 3.2),
            ("BK-1002", "Garden hose 20 m", 39.00, 2.1),
            ("BK-1003", "Potting mix 25 L", 12.50, 11.0),
            ("BK-1004", "Secateurs", 29.95, 0.3),
            ("BK-1005", "LED desk lamp", 49.00, 1.1),
            ("BK-1006", "Cast iron pan 28 cm", 89.00, 2.9),
            ("BK-1007", "Linen tea towels (3)", 19.95, 0.25),
            ("BK-1008", "Outdoor solar lights (6)", 34.95, 0.9),
            ("BK-1009", "Cordless drill 18V", 149.00, 1.6),
            ("BK-1010", "Dog bed large", 69.00, 2.4),
            ("BK-1011", "Kitchen scales", 27.50, 0.8),
            ("BK-1012", "Wheelbarrow 100 L", 129.00, 14.5)]
prod = pd.DataFrame(PRODUCTS, columns=["product_code", "description",
                                       "price", "weight_kg"])
orders = []
start = pd.Timestamp("2026-03-01", tz="UTC")
cust_ids = cust["cust_no"].unique()
for i in range(6000):
    ts = start + pd.Timedelta(minutes=int(rng.integers(0, 61 * 24 * 60)))
    pcode, desc, price, wkg = PRODUCTS[rng.integers(0, len(PRODUCTS))]
    qty = int(rng.choice([1, 1, 1, 1, 2, 2, 3, 4]))
    orders.append({"order_id": f"W{500001 + i}",
                   "customer_id": random.choice(cust_ids),
                   "order_ts_utc": ts.strftime("%Y-%m-%dT%H:%M:%SZ"),
                   "product_code": pcode, "qty": qty,
                   "unit_price": price,
                   "ship_weight": round(wkg * qty, 2),
                   "weight_unit": "kg"})
wo = pd.DataFrame(orders).sort_values("order_ts_utc").reset_index(drop=True)
wo["order_id"] = [f"W{500001 + i}" for i in range(len(wo))]
# warehouse system reports weight in grams for some rows
g = wo.sample(frac=0.22, random_state=5).index
wo.loc[g, "ship_weight"] = (wo.loc[g, "ship_weight"] * 1000).round(0)
wo.loc[g, "weight_unit"] = "g"
# marketplace orders arrive with prices in cents
c = wo.sample(frac=0.08, random_state=6).index
wo.loc[c, "unit_price"] = (wo.loc[c, "unit_price"] * 100).round(0)
# sentinels and real outliers
s = wo.sample(9, random_state=8).index
wo.loc[s, "qty"] = 999
s2 = wo.drop(s).sample(6, random_state=9).index
wo.loc[s2, "qty"] = -1
bulk = wo.drop(s).drop(s2).sample(4, random_state=10).index
wo.loc[bulk, "qty"] = [40, 60, 120, 48]
wo.loc[bulk, "ship_weight"] = (prod.set_index("product_code")
                               .loc[wo.loc[bulk, "product_code"],
                                    "weight_kg"].values
                               * wo.loc[bulk, "qty"].values).round(2)
wo.loc[bulk, "weight_unit"] = "kg"
# accounts deleted before the export: their orders remain
gone = wo.sample(23, random_state=13).index
wo.loc[gone, "customer_id"] = [f"{N + 100 + k % 5:06d}"
                               for k in range(len(gone))]
# export retried: 37 rows appear twice
dup = wo.sample(37, random_state=12)
wo = pd.concat([wo, dup]).sort_values("order_ts_utc",
                                      kind="stable")
wo = wo.reset_index(drop=True)
wo.to_csv(RAW / "web_orders.csv", index=False)

# product table with a duplicated code (price change kept as new row)
prod_dup = pd.concat([prod, pd.DataFrame(
    [("BK-1009", "Cordless drill 18V (new model)", 159.00, 1.6),
     ("BK-1006", "Cast iron pan 28 cm", 95.00, 2.9)],
    columns=prod.columns)], ignore_index=True)
prod_dup.to_csv(RAW / "products.csv", index=False)

# ---------------------------------------------------------------
# 4. Supplier price list: semicolons, decimal commas, BOM
# ---------------------------------------------------------------
sup = pd.DataFrame({
    "sku": ["NL-2201", "NL-2202", "NL-2203", "NL-2204", "NL-2205",
            "NL-2206", "NL-2207", "NL-2208"],
    "description": ["Watering can 10 L", "Bird bath, ceramic",
                    "Garden kneeler", "Hose nozzle, 8 pattern",
                    "Plant labels (50)", "Hanging basket 35 cm",
                    "Trowel, stainless", "Compost bin 220 L"],
    "unit_price_eur": ["12,40", "48,90", "9,75", "6,20", "2,95",
                       "18,60", "11,05", "1.049,00"],
    "pack_qty": [6, 2, 10, 12, 40, 4, 12, 1]})
with open(RAW / "supplier_prices.txt", "w", encoding="utf-8-sig",
          newline="") as fh:
    sup.to_csv(fh, sep=";", index=False)

# ---------------------------------------------------------------
# 5. Monthly sales report workbook (a report, not a table)
# ---------------------------------------------------------------
months = pd.period_range("2025-07", "2026-06", freq="M")
stores = {"Metro": ["Norwood", "Marion", "Mawson Lakes", "Glenelg"],
          "Regional": ["Murray Bridge", "Mount Gambier", "Renmark"]}
opened = {"Norwood": 41275, "Marion": "2014-03-17",
          "Mawson Lakes": 43709, "Glenelg": "2019-11-02",
          "Murray Bridge": 45474, "Mount Gambier": 45474,
          "Renmark": 45474}
wb = Workbook()
ws = wb.active
ws.title = "Sales"
ws["A1"] = "Banksia Home Supply"
ws["A1"].font = Font(bold=True, size=14)
ws["A2"] = "Monthly sales by store ($), FY2025-26"
ws.append([])
ws.append(["Region", "Store", "Opened"] +
           [m.strftime("%b-%y") for m in months])
grand = np.zeros(12)
for region, names in stores.items():
    sub = np.zeros(12)
    for k, st in enumerate(names):
        base = rng.integers(60, 140) * 1000
        season = 1 + 0.25 * np.sin((np.arange(12) - 3) / 12 * 2 * np.pi)
        vals = np.round(base * season * rng.normal(1, 0.06, 12), 0)
        sub += vals
        cells = [float(v) for v in vals]
        if st == "Glenelg":
            cells[4] = f"${vals[4]:,.0f}"       # typed in as text
            cells[9] = f"{vals[9]:,.0f}"
        ws.append([region if k == 0 else None, st, opened[st]] + cells)
    ws.append([None, f"{region} subtotal", None] + [float(v) for v in sub])
    grand += sub
ws.append([None, "Grand total", None] + [float(v) for v in grand])
ws.append([])
ws.append(["Source: POS extract, run 3 July 2026. Figures exclude GST."])
for row in ws.iter_rows(min_row=5, max_row=ws.max_row, min_col=3,
                        max_col=3):
    for cell in row:
        if isinstance(cell.value, int):
            cell.number_format = "General"
wb.save(RAW / "sales_report.xlsx")

# ---------------------------------------------------------------
# 6. Reference data and the SQLite database
# ---------------------------------------------------------------
pd.DataFrame(TOWNS, columns=["suburb", "postcode"]).assign(
    state=lambda d: np.where(d["postcode"].str.startswith("0"),
                             "NT", "SA")).to_csv(
    REF / "localities.csv", index=False)
pd.DataFrame({"raw_value": ["SA", "S.A.", "South Australia", "Sth Aust",
                            "NT", "N.T.", "Northern Territory"],
              "state": ["SA", "SA", "SA", "SA", "NT", "NT", "NT"]}
             ).to_csv(REF / "state_map.csv", index=False)
# suburb corrections, as reviewed by a person (Problem 8)
pd.DataFrame(
    [("Loxtn", "Loxton"), ("Mount Baker", "Mount Barker"),
     ("Mount Gambeir", "Mount Gambier"),
     ("Murray Bdge", "Murray Bridge"),
     ("Murrray Bridge", "Murray Bridge"),
     ("Naracorte", "Naracoorte"), ("Renmrk", "Renmark"),
     ("Victor Harbour", "Victor Harbor"),
     ("Victor Hbr", "Victor Harbor")],
    columns=["found", "correct"]).to_csv(
    REF / "suburb_fixes.csv", index=False)

sw = pd.read_excel(RAW / "sales_report.xlsx", header=3)
sw = sw[sw["Store"].notna() & ~sw["Store"].astype(str)
        .str.contains("total")].drop(columns=["Region", "Opened"])
sw = sw.astype(str).replace(r"\.0$", "", regex=True)

con = sqlite3.connect("data/banksia.db")
sw.to_sql("sales_wide", con, if_exists="replace", index=False)
pd.read_csv(REF / "localities.csv", dtype=str).to_sql(
    "localities", con, if_exists="replace", index=False)
cust.to_sql("mallee_customers", con, if_exists="replace", index=False)
bkc.to_sql("banksia_customers", con, if_exists="replace", index=False)
wo.to_sql("web_orders", con, if_exists="replace", index=False)
prod_dup.to_sql("products", con, if_exists="replace", index=False)
con.close()
bkc.to_csv(RAW / "banksia_customers.csv", index=False)
print("Rows:", len(cust), "customers,", len(wo), "web orders")
